"""Bounded CSV/XLSX parsing, signed read-only preview and atomic confirmation."""

import csv
import io
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import NotRequired, TypedDict
from uuid import UUID
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from django.core import signing
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import UploadedFile
from django.db import models, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.accounts.models import User

from .models import MockTest, Question, QuestionOption, QuestionType
from .services import require_owner
from .validation import validate_answers, validate_question_data

COLUMNS = (
    "question_number",
    "phase",
    "subject",
    "question_type",
    "question_text_md",
    "question_image_url",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
    "numeric_answer",
    "numeric_tolerance",
    "positive_marks",
    "negative_marks",
    "explanation_md",
)
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
SALT = "exams.question-import.v1"


class ImportRow(TypedDict):
    row: int
    values: dict[str, str]
    parse_error: NotRequired[str]


@dataclass
class ImportPreview:
    rows: list[ImportRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    token: str = ""

    @property
    def valid(self) -> bool:
        return not self.errors


def template_bytes(format: str) -> bytes:
    if format == "csv":
        stream = io.StringIO(newline="")
        csv.writer(stream).writerow(COLUMNS)
        return stream.getvalue().encode("utf-8-sig")
    if format != "xlsx":
        raise ValidationError("Template format must be csv or xlsx.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(COLUMNS)
    sheet.freeze_panes = "A2"
    from openpyxl.utils import get_column_letter

    for index, column in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = max(18, len(column) + 2)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse_upload(upload: UploadedFile) -> list[ImportRow]:
    content = upload.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValidationError("File exceeds the 2 MiB limit.")
    name = upload.name.lower()
    try:
        if name.endswith(".csv"):
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True))
        elif name.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValidationError("Expanded XLSX exceeds the 20 MiB limit.")
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise ValidationError("Use exactly one worksheet.")
                sheet = workbook.worksheets[0]
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(COLUMNS):
                    raise ValidationError("Worksheet exceeds template dimensions or 1000 rows.")
                rows = []
                for index, cells in enumerate(sheet.iter_rows(), 1):
                    if index > MAX_ROWS + 1:
                        raise ValidationError("Maximum 1000 data rows.")
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValidationError(
                            f"Row {index}: formulas and error cells are not accepted."
                        )
                    rows.append([cell.value for cell in cells])
            finally:
                workbook.close()
        else:
            raise ValidationError("Upload a UTF-8 CSV or XLSX file.")
    except ValidationError:
        raise
    except (
        UnicodeError,
        csv.Error,
        zipfile.BadZipFile,
        OSError,
        ValueError,
        KeyError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
        TypeError,
        IndexError,
        RuntimeError,
    ) as exc:
        raise ValidationError(f"Malformed file: {exc}") from exc
    if not rows or tuple(str(v).strip() for v in rows[0]) != COLUMNS:
        raise ValidationError("Header must match the downloadable template exactly.")
    if len(rows) > MAX_ROWS + 1:
        raise ValidationError("Maximum 1000 data rows.")
    parsed = []
    for index, cells in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        if len(cells) != len(COLUMNS):
            parsed.append(
                {
                    "row": index,
                    "values": {},
                    "parse_error": f"Expected {len(COLUMNS)} columns; found {len(cells)}.",
                }
            )
            continue
        parsed.append(
            {
                "row": index,
                "values": {
                    key: str(value).strip() if value is not None else ""
                    for key, value in zip(COLUMNS, cells, strict=True)
                },
            }
        )
    if not parsed:
        raise ValidationError("No question rows found.")
    return parsed


def decimal_value(value: str, name: str, default: str | None = None) -> Decimal:
    if value == "" and default is not None:
        return Decimal(default)
    try:
        number = Decimal(value)
        if not number.is_finite():
            raise InvalidOperation
        return number
    except InvalidOperation as exc:
        raise ValidationError(f"{name} must be a finite decimal number.") from exc


def build_question(mock, row, phases):
    values = row["values"]
    try:
        number = int(values["question_number"])
        phase = phases[int(values["phase"])]
    except (ValueError, KeyError) as exc:
        raise ValidationError(
            "Question number and phase must be valid integers; phase must exist."
        ) from exc
    kind = values["question_type"]
    question = Question(
        mock_test=mock,
        phase=phase,
        subject=values["subject"],
        question_number=number,
        question_type=kind,
        question_text_md=values["question_text_md"],
        question_image_url=values["question_image_url"],
        positive_marks=decimal_value(values["positive_marks"], "positive_marks"),
        negative_marks=decimal_value(values["negative_marks"], "negative_marks"),
        correct_numeric_answer=decimal_value(values["numeric_answer"], "numeric_answer")
        if values["numeric_answer"]
        else None,
        numeric_tolerance=decimal_value(values["numeric_tolerance"], "numeric_tolerance", "0"),
        explanation_md=values["explanation_md"],
        status=Question.Status.READY,
    )
    rule = next(
        (
            r
            for r in phase.scheme_phase.rules.all()
            if r.subject == question.subject and r.question_type == kind
        ),
        None,
    )
    errors = validate_question_data(question, rule)
    if errors:
        raise ValidationError(errors)
    options = []
    if kind == QuestionType.MCQ_SINGLE:
        if values["correct_option"] not in {"A", "B", "C", "D"}:
            raise ValidationError("correct_option must be exactly one label A, B, C or D.")
        for order, label in enumerate("ABCD", 1):
            option = QuestionOption(
                question=question,
                label=label,
                order=order,
                option_text_md=values[f"option_{label.lower()}"],
                is_correct=label == values["correct_option"],
            )
            option.clean_fields(exclude=["question"])
            option.clean()
            options.append(option)
    elif values["correct_option"] or any(values[f"option_{label}"] for label in "abcd"):
        raise ValidationError("Numerical rows must leave all option columns/correct_option empty.")
    errors = validate_answers(question, options)
    if errors:
        raise ValidationError(errors)
    return question, options


def validate_rows(mock, rows):
    preview = ImportPreview(rows=rows)
    if mock.status != MockTest.Status.DRAFT:
        preview.errors.append("Imports are allowed only for DRAFT mocks.")
        return preview, []
    phases = {
        p.order: p
        for p in mock.phases.select_related("scheme_phase", "mock_test").prefetch_related(
            "scheme_phase__rules"
        )
    }
    expected = {
        (r.phase_id, r.subject, r.question_type): r.question_count
        for p in mock.exam_scheme.phases.prefetch_related("rules")
        for r in p.rules.all()
    }
    existing = list(mock.questions.select_related("phase").all())
    counts = Counter((q.phase.scheme_phase_id, q.subject, q.question_type) for q in existing)
    numbers = {q.question_number for q in existing}
    contents = {(q.subject, q.question_text_md.strip(), q.question_image_url) for q in existing}
    built = []
    for row in rows:
        try:
            if row.get("parse_error"):
                raise ValidationError(row["parse_error"])
            q, options = build_question(mock, row, phases)
            if q.question_number in numbers:
                raise ValidationError("Duplicate question number in batch or existing paper.")
            if q.question_number > mock.exam_scheme.total_question_count:
                raise ValidationError("Question number exceeds the scheme total.")
            content = (q.subject, q.question_text_md.strip(), q.question_image_url)
            if content in contents:
                raise ValidationError("Duplicate question content in batch or existing paper.")
            numbers.add(q.question_number)
            contents.add(content)
            key = (q.phase.scheme_phase_id, q.subject, q.question_type)
            counts[key] += 1
            if counts[key] > expected.get(key, 0):
                raise ValidationError("Subject/type count exceeds the scheme rule.")
            built.append((q, options))
        except ValidationError as exc:
            preview.errors.extend(f"Row {row['row']}: {message}" for message in exc.messages)
    if counts != Counter(expected):
        preview.warnings.append(
            "Paper is incomplete. Final paper validation must pass before registration/scheduling."
        )
    return preview, built


def preview_import(mock_id: UUID, upload: UploadedFile, *, actor: User) -> ImportPreview:
    require_owner(actor)
    mock = MockTest.objects.select_related("exam_scheme").get(pk=mock_id)
    try:
        rows = parse_upload(upload)
    except ValidationError as exc:
        return ImportPreview(errors=exc.messages)
    preview, _ = validate_rows(mock, rows)
    if preview.valid:
        preview.token = signing.dumps(
            {"mock": str(mock.pk), "actor": str(actor.pk), "rows": rows}, salt=SALT, compress=True
        )
    return preview


@transaction.atomic
def commit_import(mock_id: UUID, token: str, *, actor: User) -> int:
    require_owner(actor)
    try:
        payload = signing.loads(token, salt=SALT, max_age=1800)
    except signing.BadSignature as exc:
        raise ValidationError("Preview is invalid/expired. Upload and preview again.") from exc
    if payload["mock"] != str(mock_id) or payload["actor"] != str(actor.pk):
        raise ValidationError("Preview belongs to a different mock or owner.")
    mock = MockTest.objects.select_for_update().select_related("exam_scheme").get(pk=mock_id)
    preview, built = validate_rows(mock, payload["rows"])
    if preview.errors:
        raise ValidationError(preview.errors)
    # Private, validated bulk insertion: the mock lock is held, every row is validated,
    # and database constraints remain active. Public bulk methods intentionally reject writes.
    models.QuerySet.bulk_create(Question.objects.all(), [q for q, _ in built])
    _insert_options([option for _, options in built for option in options])
    return len(built)


def _insert_options(options):
    models.QuerySet.bulk_create(QuestionOption.objects.all(), options)
